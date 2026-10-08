#!/usr/bin/env python3
"""Generate TrainPulse.pptx — a branded 10-slide deck (16:9)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---- Brand palette ----
BLACK  = RGBColor(0x07, 0x0D, 0x18)
NAVY   = RGBColor(0x0B, 0x1B, 0x2E)
CYAN   = RGBColor(0x22, 0xD3, 0xEE)
AMBER  = RGBColor(0xFF, 0xB0, 0x20)
CLOUD  = RGBColor(0xF1, 0xF5, 0xF9)
SLATE  = RGBColor(0x64, 0x74, 0x8B)
MUTED  = RGBColor(0x9F, 0xB0, 0xC2)
CARD   = RGBColor(0x0F, 0x24, 0x3A)
OK     = RGBColor(0x22, 0xC5, 0x5E)
DELAY  = RGBColor(0xF5, 0x9E, 0x0B)
BAD    = RGBColor(0xEF, 0x44, 0x44)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height


def bg(slide, color=BLACK):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def no_line(shape):
    shape.line.fill.background()


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def rounded(slide, l, t, w, h, color=CARD, line=None):
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    fill(sp, color)
    if line:
        sp.line.color.rgb = line
        sp.line.width = Pt(1)
    else:
        no_line(sp)
    sp.shadow.inherit = False
    return sp


def txt(slide, l, t, w, h, runs, size=18, color=CLOUD, bold=False,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font="Segoe UI",
        line_spacing=1.0, space_after=0):
    """runs: str, or list of (text,color,bold[,size]) tuples, or list of paragraphs(list of runs)."""
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

    if isinstance(runs, str):
        runs = [[(runs, color, bold, size)]]
    elif runs and isinstance(runs[0], tuple):
        runs = [runs]  # single paragraph of runs

    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        if space_after:
            p.space_after = Pt(space_after)
        for item in para:
            rtext = item[0]
            rcolor = item[1] if len(item) > 1 and item[1] else color
            rbold = item[2] if len(item) > 2 else bold
            rsize = item[3] if len(item) > 3 else size
            r = p.add_run()
            r.text = rtext
            r.font.size = Pt(rsize)
            r.font.bold = rbold
            r.font.color.rgb = rcolor
            r.font.name = font
    return tb


def accent_bar(slide):
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(0.9), Inches(1.55), Inches(0.28), Inches(0.28))
    fill(sp, AMBER); no_line(sp); sp.shadow.inherit = False


def kicker(slide, text):
    txt(slide, Inches(0.9), Inches(0.7), Inches(10), Inches(0.4),
        [(text.upper(), CYAN, True, 13)], font="Segoe UI")


def heading(slide, text, bar=True):
    if bar:
        accent_bar(slide)
    txt(slide, Inches(1.35), Inches(1.3), Inches(11), Inches(0.9),
        [(text, CLOUD, True, 30)])


def footer(slide, num):
    txt(slide, Inches(0.9), Inches(7.0), Inches(5), Inches(0.3),
        [("TRAINPULSE", SLATE, False, 10)])
    txt(slide, Inches(10.3), Inches(7.0), Inches(2.1), Inches(0.3),
        [(f"{num:02d}", CYAN, True, 11), (" / 10", SLATE, False, 11)],
        align=PP_ALIGN.RIGHT)


def logo(slide, l, t):
    # pulse line
    import math
    pts = [(0,0.55),(0.18,0.55),(0.30,0.18),(0.46,0.95),(0.60,0.08),(0.72,0.5),(1.0,0.5)]
    scale_w, scale_h = Inches(1.0), Inches(0.6)
    fs = slide.shapes.build_freeform
    # simpler: draw a connector-ish line via several small lines
    base_l, base_t = l, t
    for i in range(len(pts)-1):
        x1 = base_l + Emu(int(pts[i][0]*scale_w))
        y1 = base_t + Emu(int(pts[i][1]*scale_h))
        x2 = base_l + Emu(int(pts[i+1][0]*scale_w))
        y2 = base_t + Emu(int(pts[i+1][1]*scale_h))
        ln = slide.shapes.add_connector(2, x1, y1, x2, y2)
        ln.line.color.rgb = CYAN
        ln.line.width = Pt(3)
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL,
            base_l + Emu(int(0.80*scale_w)), base_t + Emu(int(0.30*scale_h)),
            Inches(0.16), Inches(0.16))
    fill(dot, AMBER); no_line(dot); dot.shadow.inherit = False


def chips(slide, items, top, color=CYAN):
    x = Inches(0.9)
    for label in items:
        w = Inches(0.42 + 0.105 * len(label))
        c = rounded(slide, x, top, w, Inches(0.42), CARD, line=color)
        txt(slide, x, top, w, Inches(0.42), [(label, CLOUD, True, 12)],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        x = x + w + Inches(0.18)


def card_grid(slide, cards, cols, top, bottom=Inches(6.7),
              gap=Inches(0.3), left=Inches(0.9), right=Inches(0.9)):
    rows = (len(cards) + cols - 1) // cols
    total_w = SW - left - right
    cw = Emu(int((total_w - gap * (cols - 1)) / cols))
    total_h = bottom - top
    ch = Emu(int((total_h - gap * (rows - 1)) / rows))
    for i, (title, body, amber) in enumerate(cards):
        r, c = divmod(i, cols)
        x = left + Emu(int((cw + gap) * c))
        y = top + Emu(int((ch + gap) * r))
        rounded(slide, x, y, cw, ch, CARD, line=(AMBER if amber else CYAN))
        pad = Inches(0.22)
        txt(slide, x + pad, y + pad, cw - 2*pad, Inches(0.5),
            [(title, (AMBER if amber else CYAN), True, 15)])
        txt(slide, x + pad, y + pad + Inches(0.52), cw - 2*pad, ch - Inches(0.9),
            [(body, MUTED, False, 12)], line_spacing=1.12)


# ============ SLIDE 1 — TITLE ============
s = prs.slides.add_slide(BLANK); bg(s)
logo(s, Inches(0.9), Inches(0.75))
txt(s, Inches(2.05), Inches(0.78), Inches(6), Inches(0.6),
    [("Train", CLOUD, True, 22), ("Pulse", CYAN, True, 22)], anchor=MSO_ANCHOR.MIDDLE)
txt(s, Inches(0.9), Inches(2.5), Inches(10), Inches(0.4),
    [("REAL-TIME RAILWAY TRACKING", CYAN, True, 15)])
txt(s, Inches(0.85), Inches(2.95), Inches(11.5), Inches(1.6),
    [("Train", CLOUD, True, 78), ("Pulse", CYAN, True, 78)])
txt(s, Inches(0.9), Inches(4.6), Inches(10.5), Inches(1.2),
    [("A cross-platform, map-first live train tracking companion for Indian "
      "Railways — strict TypeScript, Expo / React Native Web, and a secure Hono "
      "backend gateway.", MUTED, False, 18)], line_spacing=1.3)
chips(s, ["Android", "iOS", "Web", "TypeScript (strict)", "Truthful Telemetry"], Inches(6.0))
txt(s, Inches(0.9), Inches(7.0), Inches(6), Inches(0.3),
    [("PROJECT ANALYSIS & OVERVIEW", SLATE, False, 10)])
footer_num = s  # keep handle; footer page below
txt(s, Inches(10.3), Inches(7.0), Inches(2.1), Inches(0.3),
    [("01", CYAN, True, 11), (" / 10", SLATE, False, 11)], align=PP_ALIGN.RIGHT)

# ============ SLIDE 2 — PROBLEM ============
s = prs.slides.add_slide(BLANK); bg(s)
kicker(s, "The Problem")
heading(s, "Railway apps are cluttered & untrustworthy")
txt(s, Inches(0.9), Inches(2.0), Inches(11), Inches(0.8),
    [("Most apps bury the one thing passengers want — where is my train right "
      "now? — under ticketing, PNR, ads, and plain text lists.", MUTED, False, 16)],
    line_spacing=1.25)
card_grid(s, [
    ("Feature bloat", "Ticketing, PNR, seat maps and ads crowd out the core live-tracking use case.", False),
    ("No spatial awareness", "Location shown as textual lists instead of an intuitive map-based view.", False),
    ("Fabricated data", "Estimates presented as fact — users can't tell a real GPS fix from a guess.", True),
], 3, Inches(3.0), bottom=Inches(6.5))
footer(s, 2)

# ============ SLIDE 3 — SOLUTION ============
s = prs.slides.add_slide(BLANK); bg(s)
kicker(s, "The Solution")
heading(s, "A map-first, truthful live companion")
bullets = [
    [("A ", MUTED,0,16),("map is the primary interface", CLOUD,1,16),(" — search overlaid on a live route map.", MUTED,0,16)],
    [("Truthful telemetry: ", AMBER,1,16),("every position tagged ACTUAL_GPS vs ESTIMATED_SCHEDULE — never synthesized.", MUTED,0,16)],
    [("One codebase → ", MUTED,0,16),("Android, iOS & Web parity", CLOUD,1,16),(" via Expo + React Native Web.", MUTED,0,16)],
    [("Backend caching & coalescing", CLOUD,1,16),(" protect upstream quotas and keep loads sub-second.", MUTED,0,16)],
    [("System-aware Dark / Light theming", CLOUD,1,16),(" using official TrainPulse brand colors.", MUTED,0,16)],
]
y = Inches(2.3)
for i, b in enumerate(bullets):
    amber = (i == 1)
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.95), y + Inches(0.1),
                             Inches(0.16), Inches(0.16))
    fill(dot, AMBER if amber else CYAN); no_line(dot); dot.shadow.inherit = False
    txt(s, Inches(1.35), y, Inches(10.8), Inches(0.7), [b], line_spacing=1.15)
    y = y + Inches(0.78)
footer(s, 3)

# ============ SLIDE 4 — CORE FEATURES ============
s = prs.slides.add_slide(BLANK); bg(s)
kicker(s, "Core Features")
heading(s, "What TrainPulse does")
card_grid(s, [
    ("Journey Search", "Where-to-where search between any two stations with debounced autocomplete.", False),
    ("Live Telemetry", "Status, speed, location, next halt, arrival/departure times & delay minutes.", False),
    ("Interactive Map", "OpenStreetMap + Leaflet route map on Web & Mobile via WebView.", False),
    ("Platform Status", "Confirmed / Expected / Unavailable states with freshness timestamps.", True),
    ("Coach & My Coach", "Visual rake layout with a My Coach selector and offline persistence.", False),
    ("Zero Fabrication", "Missing GPS is explicitly flagged, never invented.", True),
], 3, Inches(2.3), bottom=Inches(6.5))
footer(s, 4)

# ============ SLIDE 5 — ARCHITECTURE ============
s = prs.slides.add_slide(BLANK); bg(s)
kicker(s, "Architecture")
heading(s, "Decoupled client to secure gateway")
# flow nodes
nodes = [
    ("Client", "Expo / RN Web + Zustand", False),
    ("Hono Gateway", "/api/v1 - Zod validation", False),
    ("Cache + Coalescer", "30-45s TTL - single-flight", False),
    ("RailRadar", "GPS telemetry (Bearer auth)", True),
]
x = Inches(0.9); y = Inches(2.15); nw = Inches(2.75); nh = Inches(1.0)
for i, (t, sub, up) in enumerate(nodes):
    rounded(s, x, y, nw, nh, CARD, line=(AMBER if up else CYAN))
    txt(s, x + Inches(0.15), y + Inches(0.15), nw - Inches(0.3), Inches(0.4),
        [(t, CLOUD, True, 14)])
    txt(s, x + Inches(0.15), y + Inches(0.55), nw - Inches(0.3), Inches(0.4),
        [(sub, SLATE, False, 10)])
    if i < 3:
        ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + nw + Inches(0.02),
                                y + Inches(0.37), Inches(0.26), Inches(0.26))
        fill(ar, CYAN); no_line(ar); ar.shadow.inherit = False
    x = x + nw + Inches(0.33)
card_grid(s, [
    ("Secret isolation", "Client bundles contain no API keys — it doesn't even know RailRadar exists.", False),
    ("Request coalescing", "50 clients polling one train = 1 outbound request, broadcast to all.", False),
    ("Graceful degradation", "On 429/503, serve stale data with X-Cache: STALE + a freshness warning.", True),
], 3, Inches(3.6), bottom=Inches(6.5))
footer(s, 5)

# ============ SLIDE 6 — TECH STACK ============
s = prs.slides.add_slide(BLANK); bg(s)
kicker(s, "Technology Stack")
heading(s, "Built on a strict, modern toolchain")
rows = [
    ("Mobile & Web UI", "React Native - Expo SDK 57 - Expo Router - React Native Web"),
    ("Language", "TypeScript — strict mode across all packages"),
    ("State & Storage", "Zustand - AsyncStorage (offline persistence)"),
    ("Mapping", "Leaflet.js - OpenStreetMap - CartoDB tiles - react-native-webview"),
    ("Validation", "Zod schemas for runtime request/response verification"),
    ("Backend", "Node.js - Hono framework - @hono/node-server"),
    ("Testing", "Vitest — route handlers, transformers, cache, schemas"),
]
y = Inches(2.15); rh = Inches(0.62)
for i, (layer, tech) in enumerate(rows):
    if i % 2 == 0:
        rounded(s, Inches(0.9), y, Inches(11.5), rh, CARD)
    txt(s, Inches(1.1), y, Inches(3.0), rh, [(layer, CYAN, True, 14)],
        anchor=MSO_ANCHOR.MIDDLE)
    txt(s, Inches(4.2), y, Inches(8.0), rh, [(tech, CLOUD, False, 13)],
        anchor=MSO_ANCHOR.MIDDLE)
    y = y + rh + Inches(0.04)
footer(s, 6)

# ============ SLIDE 7 — API ============
s = prs.slides.add_slide(BLANK); bg(s)
kicker(s, "API Surface")
heading(s, "Versioned REST endpoints  /api/v1")
api = [
    ("GET /trains/:n/live", "Live status, speed, location, next station & delay for a 5-digit train.", False),
    ("GET /trains/between", "Direct trains running between a from and to station code.", False),
    ("GET /stations/search", "Station autocomplete by partial code or name (min 2 chars).", False),
    ("GET /trains/:n/coaches", "Coach composition / rake layout for a train.", False),
    ("GET /trains/:n/platform", "Platform status: Confirmed / Expected / Unavailable.", True),
    ("STATUS MODEL", "", False),
]
card_grid(s, api[:5] + [("", "", False)], 3, Inches(2.3), bottom=Inches(6.0))
# status pills in last cell area
txt(s, Inches(8.9), Inches(4.35), Inches(3.3), Inches(0.4), [("Status model", SLATE, True, 12)])
pills = [("On time", OK), ("Delayed", DELAY), ("Disrupted", BAD)]
py = Inches(4.75)
for label, col in pills:
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(8.9), py + Inches(0.06), Inches(0.14), Inches(0.14))
    fill(dot, col); no_line(dot); dot.shadow.inherit = False
    txt(s, Inches(9.15), py, Inches(3), Inches(0.3), [(label, CLOUD, False, 13)])
    py = py + Inches(0.4)
footer(s, 7)

# ============ SLIDE 8 — BY THE NUMBERS ============
s = prs.slides.add_slide(BLANK); bg(s)
kicker(s, "By The Numbers")
heading(s, "Engineering snapshot")
stats = [
    ("3", "Platforms from one codebase", False),
    ("225+", "Automated tests (FE + BE)", False),
    ("30-45s", "Adaptive polling & cache TTL", True),
    ("0", "Secrets in the client bundle", False),
    ("71", "TypeScript source files", False),
    ("5", "Core REST endpoints", False),
    ("100%", "Expo Go compatibility", True),
    ("<1s", "Cached initial load target", False),
]
cols = 4; left = Inches(0.9); gap = Inches(0.3); top = Inches(2.4); bottom = Inches(6.4)
total_w = SW - left*2
cw = Emu(int((total_w - gap*(cols-1))/cols))
rows_n = 2
ch = Emu(int((bottom - top - gap)/rows_n))
for i, (num, lbl, amber) in enumerate(stats):
    r, c = divmod(i, cols)
    x = left + Emu(int((cw+gap)*c)); y = top + Emu(int((ch+gap)*r))
    rounded(s, x, y, cw, ch, CARD, line=(AMBER if amber else CYAN))
    txt(s, x, y + Inches(0.35), cw, Inches(0.9),
        [(num, AMBER if amber else CYAN, True, 40)], align=PP_ALIGN.CENTER)
    txt(s, x + Inches(0.12), y + Inches(1.3), cw - Inches(0.24), Inches(0.7),
        [(lbl, MUTED, False, 12)], align=PP_ALIGN.CENTER, line_spacing=1.1)
footer(s, 8)

# ============ SLIDE 9 — PRINCIPLES ============
s = prs.slides.add_slide(BLANK); bg(s)
kicker(s, "Principles")
heading(s, "Security & design discipline")
pr_bul = [
    [("Zero-trust frontend: ", CLOUD,1,16),("no third-party railway keys ever reach the client.", MUTED,0,16)],
    [("Error sanitization: ", CLOUD,1,16),("auth headers, tokens & upstream keys redacted before logging or responding.", MUTED,0,16)],
    [("Truthful by default: ", AMBER,1,16),("approximations are labeled, GPS gaps are flagged — never faked.", MUTED,0,16)],
    [("Lifecycle-aware polling: ", CLOUD,1,16),("pauses on background / hidden tab via AppState & visibilitychange.", MUTED,0,16)],
    [("Deliberately out of scope: ", CLOUD,1,16),("no ticketing, PNR, seat availability or accounts.", MUTED,0,16)],
]
y = Inches(2.3)
for i, b in enumerate(pr_bul):
    amber = (i == 2)
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.95), y + Inches(0.1), Inches(0.16), Inches(0.16))
    fill(dot, AMBER if amber else CYAN); no_line(dot); dot.shadow.inherit = False
    txt(s, Inches(1.35), y, Inches(10.8), Inches(0.7), [b], line_spacing=1.15)
    y = y + Inches(0.82)
footer(s, 9)

# ============ SLIDE 10 — CLOSE ============
s = prs.slides.add_slide(BLANK); bg(s)
logo(s, Inches(0.9), Inches(0.75))
txt(s, Inches(2.05), Inches(0.78), Inches(6), Inches(0.6),
    [("Train", CLOUD, True, 22), ("Pulse", CYAN, True, 22)], anchor=MSO_ANCHOR.MIDDLE)
txt(s, Inches(0.9), Inches(2.6), Inches(10), Inches(0.4),
    [("SUMMARY", CYAN, True, 15)])
txt(s, Inches(0.85), Inches(3.05), Inches(11.5), Inches(1.6),
    [[("The train tracker you can", CLOUD, True, 52)],
     [("actually trust.", CYAN, True, 52)]], line_spacing=1.02)
txt(s, Inches(0.9), Inches(5.1), Inches(10.5), Inches(0.8),
    [("Map-first. Cross-platform. Strictly typed. Secure by design — and honest "
      "about what it does and doesn't know.", MUTED, False, 18)], line_spacing=1.3)
chips(s, ["Expo + RN Web", "Hono Backend", "Zod DTOs", "Leaflet / OSM", "MIT License"], Inches(6.1))
txt(s, Inches(0.9), Inches(7.0), Inches(6), Inches(0.3), [("THANK YOU", SLATE, False, 10)])
txt(s, Inches(10.3), Inches(7.0), Inches(2.1), Inches(0.3),
    [("10", CYAN, True, 11), (" / 10", SLATE, False, 11)], align=PP_ALIGN.RIGHT)

out = "/projects/sandbox/TrainPulse/presentation/TrainPulse.pptx"
prs.save(out)
print("Saved:", out, "slides:", len(prs.slides._sldIdLst))
